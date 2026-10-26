"""Explicit layout for MCC fees workbook. Pending summaries are never imported."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import openpyxl

def money(value):
    if value is None or str(value).strip()=='': return None
    text=str(value).strip().lower()
    if text in ('nil','nill'): return Decimal(0)
    if not re.fullmatch(r'\d+(?:\.\d+)?(?:\s*\+\s*\d+(?:\.\d+)?)*',text):
        raise ValueError(f'Unrecognized amount: {text[:60]}')
    return sum((Decimal(x.strip()) for x in text.split('+')),Decimal(0))

def parse(data):
    w=openpyxl.load_workbook(BytesIO(data),data_only=False)
    # sheet, start/end rows, specialization, status, second-semester fee column
    layout={
      'BCA':[(4,28,'AI','active',8),(32,43,'Cyber Security','active',8),(46,48,'Without add-on','no-addon',None),(52,54,'','discontinued',None)],
      'B.com':[(4,6,'Aviation','active',8),(9,12,'CMA','active',8),(15,17,'Logistics','active',8),(29,35,'Package','active',None),(47,55,'College fees','active',None)],
      'BBA':[(4,13,'Logistics','active',9),(18,26,'Aviation','active',9),(29,30,'Hospital and Health Care','active',9),(33,36,'Unspecified — verify','active',None)],
      'FASHION DESIGNING':[(3,14,'','active',None)],
      'BSW':[(3,9,'','active',None)],'MSW':[(3,10,'','active',None)],'MSC PSYCOLOGY':[(3,26,'','active',None)],'BSC PHYCOLOGY':[(3,14,'','active',None)]}
    missing=set(layout)-set(w.sheetnames)
    if missing: raise ValueError('This import expects the MCC workbook with all eight course sheets.')
    records=[]
    for sheet,sections in layout.items():
      s=w[sheet]
      for start,end,special,status,col2 in sections:
       for row in range(start,end+1):
        name=s.cell(row,1 if sheet=='BSC PHYCOLOGY' else 2).value
        if not name: continue
        rec={'key':f'{sheet}:{row}','name':str(name).strip(),'course':sheet,'specialization':special,'status':status,'fees':[],'warnings':[],'source':f'{sheet} row {row}'}
        if 'discontinued' in str(s.cell(row,7).value).lower(): rec['status']='discontinued'
        for sem,col in [(1,3),(2,col2)]:
         if not col or sheet in ['BSW','MSW','MSC PSYCOLOGY','BSC PHYCOLOGY'] or status=='discontinued': continue
         try:
          amt=money(s.cell(row,col).value)
          college=money(s.cell(row,col+1).value)
          ours=money(s.cell(row,col+2).value)
          stated=money(s.cell(row,col+3).value)
         except ValueError as e:
          rec['warnings'].append(str(e));continue
         if amt is None:
          if college is not None or ours is not None: rec['warnings'].append(f'Sem {sem}: collection exists but fee is missing; not imported.')
          continue
         calc=amt-(college or 0)-(ours or 0)
         if stated is not None and stated!=calc: rec['warnings'].append(f'Sem {sem}: stated balance {stated}, calculated {calc}.')
         if college is None or ours is None: rec['warnings'].append(f'Sem {sem}: blank collection treated as no recorded payment; verify.')
         if calc<0: rec['warnings'].append(f'Sem {sem}: advance / excess payment {abs(calc)}.')
         rec['fees'].append({'semester':sem,'amount':str(amt),'college':str(college or 0),'our':str(ours or 0),'balance':str(calc)})
        if not rec['fees']: rec['warnings'].append('Student only: fee setup or collection information needs review.')
        records.append(rec)
    return records
